import pandas as pd
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import re

# ── Source data ────────────────────────────────────────────────────────────
PATH = '/Users/anazshazlanazizan/Downloads/Deposit Calculator Survey Responses.xlsx'
all_sheets = pd.read_excel(PATH, sheet_name=None)

df1 = all_sheets['Sheet1']
df1['ts'] = pd.to_datetime(df1['Timestamp'], utc=True)
df1['date'] = df1['ts'].dt.strftime('%Y-%m-%d')
iso = re.compile(r'^[A-Z]{2}$')
may56 = df1[df1['date'].isin(['2026-05-05', '2026-05-06'])].copy()
may56 = may56[may56['Jurisdiction'].apply(lambda x: bool(iso.match(str(x))) if pd.notna(x) else False)]
may56 = may56[may56['Respondent Type'].isin(['DI', 'NDI'])].copy()
n = len(may56)  # 30
di  = may56[may56['Respondent Type'] == 'DI']
ndi = may56[may56['Respondent Type'] == 'NDI']
n_di, n_ndi = len(di), len(ndi)

df2 = all_sheets['Interactions']
df2['ts'] = pd.to_datetime(df2['Timestamp'], format='%Y%m%d-%H%M', errors='coerce')
df2['date'] = df2['ts'].dt.strftime('%Y-%m-%d')
int_may5 = len(df2[df2['date'] == '2026-05-05'])
int_may6 = len(df2[df2['date'] == '2026-05-06'])
int_total = int_may5 + int_may6

resp_type  = may56['Respondent Type'].value_counts()
age_dist   = may56['Age Group'].value_counts().sort_index()
juris_dist = may56['Jurisdiction'].value_counts()

def multi_counts(df, col):
    c = Counter()
    for val in df[col].dropna():
        for item in str(val).split(','):
            c[item.strip()] += 1
    return c

q1_counts = multi_counts(may56, 'Q1 Trigger')
di_q1     = multi_counts(di,    'Q1 Trigger')
ndi_q1    = multi_counts(ndi,   'Q1 Trigger')

q2_counts = may56['Q2 Destination'].value_counts()
di_q2     = di['Q2 Destination'].value_counts()
ndi_q2    = ndi['Q2 Destination'].value_counts()

q3_counts = multi_counts(may56, 'Q3 Reassurance')
di_q3     = multi_counts(di,    'Q3 Reassurance')
ndi_q3    = multi_counts(ndi,   'Q3 Reassurance')

# ── Colour constants ────────────────────────────────────────────────────────
NAVY    = 'FF1F4E79'
BLUE    = 'FF2E75B6'
LBLUE   = 'FFBDD7EE'
ALTBLUE = 'FFE9F3FB'
WHITE   = 'FFFFFFFF'
YELLOW  = 'FFFFF2CC'
DARK    = 'FF1A1A1A'
GREY6   = 'FF666666'

def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)

def bdr():
    s = Side(style='thin', color='FFB8B8B8')
    return Border(left=s, right=s, top=s, bottom=s)

def fmt_title(ws, r, c, txt, c2, sz=13):
    cell = ws.cell(r, c, txt)
    cell.font = Font(name='Arial', bold=True, size=sz, color=WHITE)
    cell.fill = fill(NAVY)
    cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c2)
    ws.row_dimensions[r].height = 28

def fmt_section(ws, r, c, txt, c2):
    cell = ws.cell(r, c, txt)
    cell.font = Font(name='Arial', bold=True, size=11, color=WHITE)
    cell.fill = fill(BLUE)
    cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c2)
    ws.row_dimensions[r].height = 20

def fmt_colhdr(ws, r, c, txt, ctr=False):
    cell = ws.cell(r, c, txt)
    cell.font = Font(name='Arial', bold=True, size=10, color=DARK)
    cell.fill = fill(LBLUE)
    cell.alignment = Alignment(horizontal='center' if ctr else 'left',
                               vertical='center', wrap_text=True)
    cell.border = bdr()
    ws.row_dimensions[r].height = 18

def fmt_data(ws, r, c, val, alt=False, bold=False, ctr=False, num_fmt=None):
    cell = ws.cell(r, c, val)
    cell.font = Font(name='Arial', size=10, bold=bold, color=DARK)
    cell.fill = fill(ALTBLUE if alt else WHITE)
    cell.alignment = Alignment(horizontal='center' if ctr else 'left',
                               vertical='center', wrap_text=True)
    cell.border = bdr()
    if num_fmt:
        cell.number_format = num_fmt

def fmt_insight(ws, r, c, txt, c2):
    cell = ws.cell(r, c, txt)
    cell.font = Font(name='Arial', size=10, italic=True, color=DARK)
    cell.fill = fill(YELLOW)
    cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c2)
    ws.row_dimensions[r].height = 38

# ── Reference data ─────────────────────────────────────────────────────────
AGE  = [('a', '19 years old and below'), ('b', '20 – 29 years old'),
        ('c', '30 – 39 years old'),      ('d', '40 – 49 years old'),
        ('e', '50 – 59 years old'),      ('f', '60 years old and above')]
RESP = [('DI', 'Deposit Insurer'), ('NDI', 'Non-Deposit Insurer')]
Q1   = [('a', 'Mobile or online banking disruptions'),
        ('b', 'News of financial losses or signs of distress of a bank'),
        ('c', 'Cases of bank failures in other jurisdictions'),
        ('d', 'Negative social media rumors'),
        ('e', 'Long waiting times or queues at bank branches'),
        ('f', 'None of the above')]
Q2   = [('a', 'Another domestic bank protected by a deposit insurance scheme'),
        ('b', 'Foreign bank with a higher coverage limit where depositors already have an overseas account'),
        ('c', 'Cash at home'),
        ('d', 'E-wallets'),
        ('e', 'Investment (e.g. gold, investment account, crypto, unit trust)')]
Q3   = [('a', 'Clear and regular public statements issued by the affected bank'),
        ('b', 'Statements or official communication from regulators'),
        ('c', 'Mainstream media clarification'),
        ('d', 'Social media posting by key opinion leaders or influencers'),
        ('e', 'Seeing others remain calm'),
        ('f', 'Deposits are fully protected by a deposit insurance scheme')]

wb = Workbook()

# ══════════════════════════════════════════════════════════════════════════
# SHEET 1 — Reference
# ══════════════════════════════════════════════════════════════════════════
ws_ref = wb.active
ws_ref.title = 'Reference'
ws_ref.sheet_view.showGridLines = False
ws_ref.column_dimensions['A'].width = 10
ws_ref.column_dimensions['B'].width = 76

r = 1
fmt_title(ws_ref, r, 1, 'Survey Answer Key  —  Category Code Reference', 2)
r += 2

REF_SECS = [
    ('Age Group',
     AGE,
     'All respondents — demographic gate question.'),
    ('Respondent Type  (Gate Question)',
     RESP,
     'Q0 — Determines track: Track A (DI) or Track B (NDI).'),
    ('Q1  —  Deposit Withdrawal Trigger',
     Q1,
     'DI (Track A): "In your jurisdiction, which events are most likely to trigger depositors to withdraw their deposits?"\n'
     'NDI (Track B): "Which events would make you consider withdrawing your deposits?"\n'
     'Multi-select.'),
    ('Q2  —  Destination for Moved Funds',
     Q2,
     'DI (Track A): "If depositors were to move their money during a period of stress, where are they most likely to place it?"\n'
     'NDI (Track B): "If you were to move your money, where would you be most likely to place it?"\n'
     'Single-select.'),
    ('Q3  —  Reassurance Factors',
     Q3,
     'DI (Track A): "Which actions or communications are most effective in reassuring depositors and preventing unnecessary withdrawals?"\n'
     'NDI (Track B): "Which of these would make you feel reassured enough to keep your deposits?"\n'
     'Multi-select.'),
]

for sec_title, items, note in REF_SECS:
    fmt_section(ws_ref, r, 1, sec_title, 2)
    r += 1
    nl = note.count('\n')
    lbl = ws_ref.cell(r, 1, 'Note:')
    lbl.font = Font(name='Arial', size=9, italic=True, color=GREY6)
    body = ws_ref.cell(r, 2, note)
    body.font = Font(name='Arial', size=9, italic=True, color=GREY6)
    body.alignment = Alignment(wrap_text=True)
    ws_ref.row_dimensions[r].height = 16 * (nl + 1) + 6
    r += 1
    fmt_colhdr(ws_ref, r, 1, 'Code', ctr=True)
    fmt_colhdr(ws_ref, r, 2, 'Answer / Label')
    r += 1
    for i, (code, label) in enumerate(items):
        alt = (i % 2 == 1)
        fmt_data(ws_ref, r, 1, code, alt=alt, ctr=True)
        fmt_data(ws_ref, r, 2, label, alt=alt)
        r += 1
    r += 1

for txt in [
    'Tool: All entries reflect the same single tool (APRC2026-DCC). Column excluded from analysis.',
    'Jurisdiction: ISO 3166-1 alpha-2 country codes (e.g. MY = Malaysia, NZ = New Zealand, ID = Indonesia).',
]:
    cell = ws_ref.cell(r, 1, txt)
    cell.font = Font(name='Arial', size=9, italic=True, color='FF999999')
    ws_ref.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1

# ══════════════════════════════════════════════════════════════════════════
# SHEET 2 — Table A: Survey Findings
# ══════════════════════════════════════════════════════════════════════════
ws_a = wb.create_sheet('Table A – Survey')
ws_a.sheet_view.showGridLines = False
ws_a.column_dimensions['A'].width = 10
ws_a.column_dimensions['B'].width = 52
ws_a.column_dimensions['C'].width = 14
ws_a.column_dimensions['D'].width = 12
ws_a.column_dimensions['E'].width = 14
ws_a.column_dimensions['F'].width = 12
NCOLS_A = 6  # all sections span A:F

r = 1
fmt_title(ws_a, r, 1, f'Survey Findings  |  May 5 – 6, 2026  (n = {n})', NCOLS_A)
r += 2

# ── Section 1: Respondent Profile (4-col data; E–F left blank) ────────────
fmt_section(ws_a, r, 1, '1.  Respondent Profile', NCOLS_A); r += 1
for c, h, ctr in [(1,'Dimension',False),(2,'Category',False),(3,'Count',True),(4,'% of n=30',True)]:
    fmt_colhdr(ws_a, r, c, h, ctr)
# blank header cells E & F to carry the border
for c in [5, 6]:
    fmt_colhdr(ws_a, r, c, '', ctr=True)
r += 1

others_count = int(juris_dist.drop(['MY','NZ','UZ','ID','IN']).sum())
profile_rows = [
    ('Respondent Type',      'Deposit Insurer (DI)',       resp_type.get('DI', 0)),
    ('',                     'Non-Deposit Insurer (NDI)',  resp_type.get('NDI', 0)),
    ('Age Group',            '20 – 29 years old',          age_dist.get('b', 0)),
    ('',                     '30 – 39 years old',          age_dist.get('c', 0)),
    ('',                     '40 – 49 years old',          age_dist.get('d', 0)),
    ('',                     '50 – 59 years old',          age_dist.get('e', 0)),
    ('Jurisdiction (Top 5)', 'Malaysia (MY)',               juris_dist.get('MY', 0)),
    ('',                     'New Zealand (NZ)',             juris_dist.get('NZ', 0)),
    ('',                     'Uzbekistan (UZ)',              juris_dist.get('UZ', 0)),
    ('',                     'Indonesia (ID)',               juris_dist.get('ID', 0)),
    ('',                     'India (IN)',                   juris_dist.get('IN', 0)),
    ('',                     'Others (5 countries)',         others_count),
]
for i, (dim, cat, cnt) in enumerate(profile_rows):
    alt = (i % 2 == 1)
    fmt_data(ws_a, r, 1, dim, alt=alt, bold=(dim != ''))
    fmt_data(ws_a, r, 2, cat, alt=alt)
    fmt_data(ws_a, r, 3, int(cnt), alt=alt, ctr=True)
    fmt_data(ws_a, r, 4, int(cnt)/n, alt=alt, ctr=True, num_fmt='0.0%')
    fmt_data(ws_a, r, 5, '', alt=alt)
    fmt_data(ws_a, r, 6, '', alt=alt)
    r += 1
r += 1

# ── Helper: sub-header row for DI / NDI group labels ─────────────────────
def di_ndi_group_headers(ws, r, n_di, n_ndi):
    fmt_colhdr(ws, r, 1, 'Code', ctr=True)
    fmt_colhdr(ws, r, 2, 'Answer', ctr=False)
    # DI group header spanning C:D
    cell = ws.cell(r, 3, f'Deposit Insurer  (n={n_di})')
    cell.font = Font(name='Arial', bold=True, size=10, color=DARK)
    cell.fill = fill(LBLUE)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = bdr()
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    # NDI group header spanning E:F
    cell2 = ws.cell(r, 5, f'Non-Deposit Insurer  (n={n_ndi})')
    cell2.font = Font(name='Arial', bold=True, size=10, color=DARK)
    cell2.fill = fill(LBLUE)
    cell2.alignment = Alignment(horizontal='center', vertical='center')
    cell2.border = bdr()
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 18

def di_ndi_col_headers(ws, r, lbl_di='Responses', lbl_ndi='Responses'):
    fmt_colhdr(ws, r, 1, '', ctr=True)
    fmt_colhdr(ws, r, 2, '', ctr=False)
    fmt_colhdr(ws, r, 3, lbl_di, ctr=True)
    fmt_colhdr(ws, r, 4, '% Selected', ctr=True)
    fmt_colhdr(ws, r, 5, lbl_ndi, ctr=True)
    fmt_colhdr(ws, r, 6, '% Selected', ctr=True)

# ── Section 2: Q1 Triggers ────────────────────────────────────────────────
fmt_section(ws_a, r, 1, '2.  Q1  —  Deposit Withdrawal Triggers  (multi-select; % within each respondent group)', NCOLS_A); r += 1
di_ndi_group_headers(ws_a, r, n_di, n_ndi); r += 1
di_ndi_col_headers(ws_a, r); r += 1

for i, (code, label) in enumerate(Q1):
    dc  = di_q1.get(code, 0)
    nc  = ndi_q1.get(code, 0)
    alt = (i % 2 == 1)
    fmt_data(ws_a, r, 1, code, alt=alt, ctr=True)
    fmt_data(ws_a, r, 2, label, alt=alt)
    fmt_data(ws_a, r, 3, dc, alt=alt, ctr=True)
    fmt_data(ws_a, r, 4, dc/n_di if n_di else 0, alt=alt, ctr=True, num_fmt='0.0%')
    fmt_data(ws_a, r, 5, nc, alt=alt, ctr=True)
    fmt_data(ws_a, r, 6, nc/n_ndi if n_ndi else 0, alt=alt, ctr=True, num_fmt='0.0%')
    r += 1
r += 1

# ── Section 3: Q2 Destination ─────────────────────────────────────────────
fmt_section(ws_a, r, 1, '3.  Q2  —  Destination for Moved Funds  (single-select; % within each respondent group)', NCOLS_A); r += 1
di_ndi_group_headers(ws_a, r, n_di, n_ndi); r += 1
di_ndi_col_headers(ws_a, r, lbl_di='Responses', lbl_ndi='Responses'); r += 1

for i, (code, label) in enumerate(Q2):
    dc  = int(di_q2.get(code, 0))
    nc  = int(ndi_q2.get(code, 0))
    alt = (i % 2 == 1)
    fmt_data(ws_a, r, 1, code, alt=alt, ctr=True)
    fmt_data(ws_a, r, 2, label, alt=alt)
    fmt_data(ws_a, r, 3, dc, alt=alt, ctr=True)
    fmt_data(ws_a, r, 4, dc/n_di if n_di else 0, alt=alt, ctr=True, num_fmt='0.0%')
    fmt_data(ws_a, r, 5, nc, alt=alt, ctr=True)
    fmt_data(ws_a, r, 6, nc/n_ndi if n_ndi else 0, alt=alt, ctr=True, num_fmt='0.0%')
    r += 1
r += 1

# ── Section 4: Q3 Reassurance ─────────────────────────────────────────────
fmt_section(ws_a, r, 1, '4.  Q3  —  Reassurance Factors  (multi-select; % within each respondent group)', NCOLS_A); r += 1
di_ndi_group_headers(ws_a, r, n_di, n_ndi); r += 1
di_ndi_col_headers(ws_a, r); r += 1

for i, (code, label) in enumerate(Q3):
    dc  = di_q3.get(code, 0)
    nc  = ndi_q3.get(code, 0)
    alt = (i % 2 == 1)
    fmt_data(ws_a, r, 1, code, alt=alt, ctr=True)
    fmt_data(ws_a, r, 2, label, alt=alt)
    fmt_data(ws_a, r, 3, dc, alt=alt, ctr=True)
    fmt_data(ws_a, r, 4, dc/n_di if n_di else 0, alt=alt, ctr=True, num_fmt='0.0%')
    fmt_data(ws_a, r, 5, nc, alt=alt, ctr=True)
    fmt_data(ws_a, r, 6, nc/n_ndi if n_ndi else 0, alt=alt, ctr=True, num_fmt='0.0%')
    r += 1
r += 1

# ── Key Insights A ────────────────────────────────────────────────────────
fmt_section(ws_a, r, 1, 'Key Insights', NCOLS_A); r += 1
for txt in [
    "1  DI respondents identify negative social media rumors (d) as the top trigger (70.4%), while zero NDI respondents selected it — suggesting DI professionals view social media contagion as a systemic depositor risk that depositors themselves may not share.",
    "2  Across both groups, another DI-protected domestic bank and cash at home dominate as fund destinations. NDI respondents show a comparatively broader spread (33.3% each across a, c, and e), though this reflects only 3 respondents and should be interpreted with caution.",
    "3  Regulator communication (b) is the top reassurance factor for both groups (DI: 70.4%, NDI: 66.7%), making it the single point of consensus. Notable DI-only signals: KOL/influencer posts (d) at 40.7% vs 0% for NDI. Deposit insurance coverage (f) and direct bank statements (a) score higher among NDI (66.7% each) than DI (~30–33%).",
    f"4  NDI sample is very small (n={n_ndi} vs DI n={n_di}). NDI percentages should be read as directional indicators only and not treated as statistically robust.",
]:
    fmt_insight(ws_a, r, 1, txt, NCOLS_A); r += 1

# ══════════════════════════════════════════════════════════════════════════
# SHEET 3 — Table B: Interactions
# ══════════════════════════════════════════════════════════════════════════
ws_b = wb.create_sheet('Table B – Interactions')
ws_b.sheet_view.showGridLines = False
ws_b.column_dimensions['A'].width = 18
ws_b.column_dimensions['B'].width = 16
ws_b.column_dimensions['C'].width = 16
ws_b.column_dimensions['D'].width = 22

r = 1
fmt_title(ws_b, r, 1, 'Tool Interaction Log  |  May 5 – 6, 2026', 4)
r += 2

fmt_section(ws_b, r, 1, 'Daily Interaction Volume', 4); r += 1
for c, h, ctr in [(1,'Date',False),(2,'Day',False),(3,'Interactions',True),(4,'% of 2-Day Total',True)]:
    fmt_colhdr(ws_b, r, c, h, ctr)
r += 1

for i, (date, day, cnt, pct) in enumerate([
    ('5 May 2026', 'Tuesday',   int_may5, int_may5/int_total),
    ('6 May 2026', 'Wednesday', int_may6, int_may6/int_total),
]):
    alt = (i % 2 == 1)
    fmt_data(ws_b, r, 1, date, alt=alt)
    fmt_data(ws_b, r, 2, day,  alt=alt)
    fmt_data(ws_b, r, 3, cnt,  alt=alt, ctr=True)
    fmt_data(ws_b, r, 4, pct,  alt=alt, ctr=True, num_fmt='0.0%')
    r += 1

for c, val, num_fmt in [(1,'Total',None),(2,'',None),(3,int_total,None),(4,1.0,'0.0%')]:
    cell = ws_b.cell(r, c, val)
    cell.font = Font(name='Arial', bold=True, size=10, color=WHITE)
    cell.fill = fill(NAVY)
    cell.alignment = Alignment(horizontal='center' if c >= 3 else 'left', vertical='center')
    cell.border = bdr()
    if num_fmt:
        cell.number_format = num_fmt
r += 2

fmt_section(ws_b, r, 1, 'Key Insights', 4); r += 1
for txt in [
    "1  92 interactions recorded over May 5–6. May 5 alone accounted for 59.8% of the 2-day volume, consistent with a high-engagement event or presentation day.",
    "2  A 33% day-on-day decline (55 → 37) from May 5 to May 6 reflects the typical post-peak pattern: strong initial engagement followed by self-directed follow-up activity.",
    "3  Context: the full log spans Apr 29 – May 6 (172 total interactions). The May 5–6 window represents 53.5% of all activity recorded to date, indicating a significant acceleration in engagement across the final two days of the period.",
]:
    fmt_insight(ws_b, r, 1, txt, 4); r += 1

# ── Save ───────────────────────────────────────────────────────────────────
OUT = '/Users/anazshazlanazizan/Downloads/APRC2026_DCC_Findings_May5-6.xlsx'
wb.save(OUT)
print(f'Saved: {OUT}')
